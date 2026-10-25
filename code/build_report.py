"""Fill the official sprint template (.docx) from a small Markdown file and export a PDF with LibreOffice.

Markdown subset: first line '# Title'; a line 'AUTHOR: name | affiliation'; a block 'ABSTRACT:' up to a blank line;
'## H2', '### H3', paragraphs, '- ' bullets, '1. ' numbered items, '![caption](path)' figures (width 6.2 in),
pipe tables with a header row ('|a|b|' then '|---|---|'), and '**bold**' / '*italic*' inline.
"""
import copy
import re
import subprocess
import sys
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE.parent / "official/template/template.docx"


CJK = re.compile(r"([\u3000-\u30ff\u3400-\u9fff\uff00-\uffef]+)")
CJK_FONT = "BIZ UDGothic"  # the template font has no Japanese glyphs


def add_plain(p, text, **fmt):
    for seg in CJK.split(text):
        if not seg:
            continue
        r = p.add_run(seg)
        for k, v in fmt.items():
            setattr(r, k, v)
        if CJK.fullmatch(seg):
            r.font.name = CJK_FONT
            r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), CJK_FONT)


def add_runs(p, text):
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            add_plain(p, part[2:-2], bold=True)
        elif part.startswith("*"):
            add_plain(p, part[1:-1], italic=True)
        elif part.startswith("`"):
            r = p.add_run(part[1:-1]); r.font.name = "Courier New"
        else:
            add_plain(p, part)


def set_borders(table):
    from docx.oxml import OxmlElement
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "999999")
        b.append(e)
    tblPr.append(b)


def set_cell_text(p_el, text):
    """Replace all text in a w:p element, keeping the first run's formatting."""
    runs = p_el.findall(qn("w:r"))
    first = None
    for r in runs:
        if r.find(qn("w:footnoteReference")) is not None:
            continue  # keep the template's title footnote
        if r.find(qn("w:t")) is not None and first is None:
            first = r
        else:
            p_el.remove(r)
    if first is None:
        return
    for t in first.findall(qn("w:t"))[1:]:
        first.remove(t)
    t = first.find(qn("w:t")); t.text = text; t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def main(md_path, out_docx):
    lines = Path(md_path).read_text().splitlines()
    title = lines[0].lstrip("# ").strip()
    authors = [l[7:].strip() for l in lines if l.startswith("AUTHOR:")]
    i_abs = next(i for i, l in enumerate(lines) if l.startswith("ABSTRACT:"))
    j = i_abs + 1
    abs_lines = []
    while j < len(lines) and lines[j].strip():
        abs_lines.append(lines[j].strip()); j += 1
    abstract = " ".join(abs_lines)
    body_lines = lines[j:]

    d = docx.Document(str(TEMPLATE))
    t0 = d.tables[0]._tbl
    # title
    for p in t0.iter(qn("w:p")):
        txt = "".join(x.text or "" for x in p.iter(qn("w:t")))
        if txt.strip() == "PROJECT TITLE":
            set_cell_text(p, title)
    # authors: nested 2x3 table; each cell has text nodes 'Author name k' and 'Affiliation'
    k = 0
    for t in list(t0.iter(qn("w:t"))):
        m = re.fullmatch(r"Author name (\d)", (t.text or "").strip())
        if m:
            k = int(m.group(1)) - 1
            t.text = authors[k].split("|")[0].strip() if k < len(authors) else ""
        elif (t.text or "").strip() == "Affiliation":
            t.text = authors[k].split("|")[1].strip() if k < len(authors) else ""
    if len(authors) == 1:
        # LibreOffice renders the template's 3x0.48in author grid unreadably narrow; for a single author,
        # remove that grid and put the author above "With / Apart Research" in the centred cell.
        from docx.oxml import OxmlElement
        for tbl in list(t0.iter(qn("w:tbl"))):
            grid = [g.get(qn("w:w")) for g in tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol"))]
            if grid == ["690", "690", "690"]:
                tc = tbl.getparent()
                tc.remove(tbl)
                if tc.find(qn("w:p")) is None:
                    tc.append(OxmlElement("w:p"))
        with_p = next(p for p in t0.iter(qn("w:p")) if "".join(x.text or "" for x in p.iter(qn("w:t"))).strip().startswith("With"))
        name, aff = [s.strip() for s in authors[0].split("|")]
        for txt in (name, aff, ""):
            np_ = copy.deepcopy(with_p)
            for br in np_.iter(qn("w:br")):
                br.getparent().remove(br)
            ts = list(np_.iter(qn("w:t")))
            for t in ts[1:]:
                t.text = ""
            if ts:
                ts[0].text = txt
            with_p.addprevious(np_)
    # abstract
    for p in t0.iter(qn("w:p")):
        txt = "".join(x.text or "" for x in p.iter(qn("w:t")))
        if txt.startswith("Summarize your project in"):
            set_cell_text(p, abstract)
    # drop the guidance table and every body element after the first table
    body = d.element.body
    started = False
    for el in list(body):
        if el is t0:
            started = True
            continue
        if started and el.tag != qn("w:sectPr"):
            body.remove(el)

    in_table = None
    made_tables = []
    for raw in body_lines:
        l = raw.rstrip()
        if l.startswith("|"):
            cells = [c.strip() for c in l.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                continue
            if in_table is None:
                in_table = d.add_table(rows=0, cols=len(cells))
                set_borders(in_table)
                made_tables.append(in_table)
            row = in_table.add_row()
            for c, txt in zip(row.cells, cells):
                c.paragraphs[0].text = ""
                add_runs(c.paragraphs[0], txt)
                for r in c.paragraphs[0].runs:
                    r.font.size = Pt(8.5)
            continue
        in_table = None
        if not l.strip():
            continue
        if l.startswith("## "):
            d.add_paragraph(l[3:], style="Heading 2")
        elif l.startswith("### "):
            d.add_paragraph(l[4:], style="Heading 3")
        elif l.startswith("!["):
            m = re.match(r"!\[(.*)\]\((.*)\)", l)
            d.add_picture(str((Path(md_path).parent / m.group(2)).resolve()), width=Inches(6.2))
            cap = d.add_paragraph()
            add_runs(cap, m.group(1))
            for r in cap.runs:
                r.font.size = Pt(9)
        elif l.startswith("- "):
            p = d.add_paragraph(style="List Bullet") if "List Bullet" in [s.name for s in d.styles] else d.add_paragraph()
            add_runs(p, l[2:] if p.style.name == "List Bullet" else "• " + l[2:])
            p.paragraph_format.space_after = Pt(2)
        elif re.match(r"\d+\. ", l):
            p = d.add_paragraph()
            add_runs(p, l)
            p.paragraph_format.space_after = Pt(2)
        else:
            p = d.add_paragraph()
            add_runs(p, l)
            p.paragraph_format.space_after = Pt(5)
    # first column wide enough for model names, the rest share the remaining text width (6.3 in)
    from docx.oxml import OxmlElement
    for tb in made_tables:
        n = len(tb.columns)
        first = 1.45 if n > 3 else 2.0
        widths = [Inches(first)] + [Inches((6.3 - first) / (n - 1))] * (n - 1)
        tb.autofit = False
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tb._tbl.tblPr.append(lay)
        for gc, w in zip(tb._tbl.tblGrid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(w.inches * 1440)))
        for row in tb.rows:
            for c, w in zip(row.cells, widths):
                c.width = w
    d.save(out_docx)
    outdir = str(Path(out_docx).parent)
    r = subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", outdir, out_docx],
                       capture_output=True, text=True, timeout=300)
    print(r.stdout.strip(), r.stderr.strip()[:300])


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
